import io
from typing import List, Dict, Any
from app.services.firestore_service import get_trips_by_owner, get_vehicles_by_owner, get_drivers_by_owner

def format_inr(amount: float) -> str:
    sign = "-" if amount < 0 else ""
    amount = abs(amount)
    s, *d = f"{amount:.2f}".split(".")
    if len(s) > 3:
        last3 = s[-3:]
        s = s[:-3]
        groups = []
        while len(s) > 2:
            groups.insert(0, s[-2:])
            s = s[:-2]
        if s:
            groups.insert(0, s)
        formatted_int = ",".join(groups) + "," + last3
    else:
        formatted_int = s
    dec = f".{d[0]}" if d else ".00"
    return f"{sign}₹{formatted_int}{dec}"

# Report Calculations
def calculate_daily_report(owner_id: str) -> Dict[str, Any]:
    trips = get_trips_by_owner(owner_id)
    daily_map = {}
    
    for t in trips:
        d = t.get("date", "Unknown")
        if d not in daily_map:
            daily_map[d] = {
                "date": d,
                "totalTrips": 0,
                "totalDistanceKm": 0.0,
                "totalEarnings": 0.0,
                "totalFuelCost": 0.0,
                "totalTollAmount": 0.0,
                "totalLoadingUnloadingAmount": 0.0,
                "totalOtherExpenses": 0.0,
                "netProfit": 0.0
            }
        item = daily_map[d]
        item["totalTrips"] += 1
        item["totalDistanceKm"] += float(t.get("distanceKm", 0.0) or 0.0)
        
        fuel = t.get("fuelCost")
        if fuel is not None:
            item["totalFuelCost"] += float(fuel)
            
        item["totalTollAmount"] += float(t.get("tollAmount", 0.0) or 0.0)
        item["totalLoadingUnloadingAmount"] += float(t.get("loadingUnloadingAmount", 0.0) or 0.0)
        item["totalOtherExpenses"] += float(t.get("otherExpenses", 0.0) or 0.0)
        
        if t.get("earnings") is not None:
            item["totalEarnings"] += float(t.get("earnings"))
        if t.get("profit") is not None:
            item["netProfit"] += float(t.get("profit"))
        
    items = sorted(list(daily_map.values()), key=lambda x: x["date"], reverse=True)
    
    grand_earnings = sum(i["totalEarnings"] for i in items)
    grand_fuel = sum(i["totalFuelCost"] for i in items)
    grand_toll = sum(i["totalTollAmount"] for i in items)
    grand_loading = sum(i["totalLoadingUnloadingAmount"] for i in items)
    grand_other = sum(i["totalOtherExpenses"] for i in items)
    grand_expenses = grand_fuel + grand_toll + grand_loading + grand_other
    grand_profit = sum(i["netProfit"] for i in items)
    
    return {
        "items": items,
        "grandTotalEarnings": grand_earnings,
        "grandTotalFuelCost": grand_fuel,
        "grandTotalExpenses": grand_expenses,
        "grandTotalProfit": grand_profit
    }

def calculate_monthly_report(owner_id: str) -> Dict[str, Any]:
    trips = get_trips_by_owner(owner_id)
    monthly_map = {}
    
    for t in trips:
        date_str = t.get("date", "")
        m = date_str[:7] if len(date_str) >= 7 else "Unknown"
        if m not in monthly_map:
            monthly_map[m] = {
                "month": m,
                "totalTrips": 0,
                "totalDistanceKm": 0.0,
                "totalEarnings": 0.0,
                "totalFuelCost": 0.0,
                "totalTollAmount": 0.0,
                "totalLoadingUnloadingAmount": 0.0,
                "totalOtherExpenses": 0.0,
                "netProfit": 0.0
            }
        item = monthly_map[m]
        item["totalTrips"] += 1
        item["totalDistanceKm"] += float(t.get("distanceKm", 0.0) or 0.0)
        
        fuel = t.get("fuelCost")
        if fuel is not None:
            item["totalFuelCost"] += float(fuel)
            
        item["totalTollAmount"] += float(t.get("tollAmount", 0.0) or 0.0)
        item["totalLoadingUnloadingAmount"] += float(t.get("loadingUnloadingAmount", 0.0) or 0.0)
        item["totalOtherExpenses"] += float(t.get("otherExpenses", 0.0) or 0.0)
        
        if t.get("earnings") is not None:
            item["totalEarnings"] += float(t.get("earnings"))
        if t.get("profit") is not None:
            item["netProfit"] += float(t.get("profit"))
        
    items = sorted(list(monthly_map.values()), key=lambda x: x["month"], reverse=True)
    
    grand_earnings = sum(i["totalEarnings"] for i in items)
    grand_fuel = sum(i["totalFuelCost"] for i in items)
    grand_toll = sum(i["totalTollAmount"] for i in items)
    grand_loading = sum(i["totalLoadingUnloadingAmount"] for i in items)
    grand_other = sum(i["totalOtherExpenses"] for i in items)
    grand_expenses = grand_fuel + grand_toll + grand_loading + grand_other
    grand_profit = sum(i["netProfit"] for i in items)
    
    return {
        "items": items,
        "grandTotalEarnings": grand_earnings,
        "grandTotalFuelCost": grand_fuel,
        "grandTotalExpenses": grand_expenses,
        "grandTotalProfit": grand_profit
    }

def calculate_profit_report(owner_id: str) -> Dict[str, Any]:
    trips = get_trips_by_owner(owner_id)
    vehicles = get_vehicles_by_owner(owner_id)
    drivers = get_drivers_by_owner(owner_id)
    
    import datetime
    today_str = datetime.date.today().isoformat()
    trips_today = sum(1 for t in trips if t.get("date") == today_str)
    
    total_trips = len(trips)
    total_dist = sum(float(t.get("distanceKm", 0.0) or 0.0) for t in trips)
    total_fuel = sum(float(t.get("fuelCost")) for t in trips if t.get("fuelCost") is not None)
    total_toll = sum(float(t.get("tollAmount", 0.0) or 0.0) for t in trips)
    total_loading = sum(float(t.get("loadingUnloadingAmount", 0.0) or 0.0) for t in trips)
    total_other = sum(float(t.get("otherExpenses", 0.0) or 0.0) for t in trips)
    total_expenses = total_fuel + total_toll + total_loading + total_other
    
    total_earnings = sum(float(t.get("earnings")) for t in trips if t.get("earnings") is not None)
    net_profit = sum(float(t.get("profit")) for t in trips if t.get("profit") is not None)
    
    profit_margin = (net_profit / total_earnings * 100.0) if total_earnings > 0 else 0.0
    
    return {
        "totalTrips": total_trips,
        "totalDistanceKm": total_dist,
        "totalEarnings": total_earnings,
        "totalFuelCost": total_fuel,
        "totalTollAmount": total_toll,
        "totalLoadingUnloadingAmount": total_loading,
        "totalOtherExpenses": total_other,
        "netProfit": net_profit,
        "profitMarginPercent": round(profit_margin, 2),
        "totalVehicles": len(vehicles),
        "activeDrivers": len(drivers),
        "tripsToday": trips_today
    }


# PDF Generator (ReportLab)
def generate_pdf_report(owner_id: str) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    
    trips = get_trips_by_owner(owner_id)
    profit_data = calculate_profit_report(owner_id)
    
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=28, leftMargin=28, topMargin=28, bottomMargin=28)
    story = []
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#2563EB"),
        spaceAfter=10
    )
    
    story.append(Paragraph("Fleet Management - Operations & Freight Economics Report (INR)", title_style))
    story.append(Spacer(1, 8))
    
    # Summary Table with INR formatting
    summary_data = [
        ["Total Trips", str(profit_data["totalTrips"]), "Total Revenue", format_inr(profit_data['totalEarnings'])],
        ["Total Fuel Cost", format_inr(profit_data['totalFuelCost']), "Total Toll Plaza", format_inr(profit_data['totalTollAmount'])],
        ["Site / Hamali Charges", format_inr(profit_data['totalLoadingUnloadingAmount']), "Other Expenses", format_inr(profit_data['totalOtherExpenses'])],
        ["Net Profit", format_inr(profit_data['netProfit']), "Profit Margin", f"{profit_data['profitMarginPercent']}%"]
    ]
    summary_table = Table(summary_data, colWidths=[130, 140, 140, 140])
    summary_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F1F5F9")),
        ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor("#1F2937")),
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 5),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
    ]))
    story.append(summary_table)
    story.append(Spacer(1, 16))
    
    story.append(Paragraph("Itemized Trips Log", styles['Heading2']))
    story.append(Spacer(1, 6))
    
    headers = ["Date", "Driver", "Vehicle", "Weight", "Revenue", "Fuel", "Toll/Site", "Profit"]
    table_data = [headers]
    
    for t in sorted(trips, key=lambda x: x.get("date", ""), reverse=True):
        weight_str = f"{float(t.get('materialWeight')):.1f}T" if t.get("materialWeight") is not None else "-"
        rev_str = format_inr(float(t.get('earnings'))) if t.get('earnings') is not None else "Pending"
        fuel_str = format_inr(float(t.get('fuelCost'))) if t.get('fuelCost') is not None else "None"
        toll_and_site = float(t.get('tollAmount', 0.0) or 0.0) + float(t.get('loadingUnloadingAmount', 0.0) or 0.0)
        profit_str = format_inr(float(t.get('profit'))) if t.get('profit') is not None else "Pending"
        
        row = [
            str(t.get("date", "")),
            str(t.get("driverId", ""))[:8],
            str(t.get("vehicleId", ""))[:8],
            weight_str,
            rev_str,
            fuel_str,
            format_inr(toll_and_site),
            profit_str
        ]
        table_data.append(row)
        
    trips_table = Table(table_data, colWidths=[65, 70, 70, 50, 75, 70, 75, 75])
    trips_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#2563EB")),
        ('TEXTCOLOR', (0,0), (-1,0), colors.white),
        ('FONTNAME', (0,0), (-1,0), 'Helvetica-Bold'),
        ('ALIGN', (3,0), (-1,-1), 'RIGHT'),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(trips_table)
    
    doc.build(story)
    pdf_bytes = buffer.getvalue()
    buffer.close()
    return pdf_bytes

# Excel Generator (OpenPyXL)
def generate_excel_report(owner_id: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    
    trips = get_trips_by_owner(owner_id)
    profit_data = calculate_profit_report(owner_id)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fleet Report"
    
    ws.merge_cells("A1:K1")
    title_cell = ws["A1"]
    title_cell.value = "Fleet Management - Operations & Freight Report (INR)"
    title_cell.font = Font(name="Arial", size=16, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35
    
    summary_headers = ["Total Trips", "Total Revenue", "Fuel Cost", "Toll Plaza", "Site Charges", "Other Expenses", "Net Profit", "Profit Margin"]
    summary_values = [
        profit_data["totalTrips"],
        format_inr(profit_data["totalEarnings"]),
        format_inr(profit_data["totalFuelCost"]),
        format_inr(profit_data["totalTollAmount"]),
        format_inr(profit_data["totalLoadingUnloadingAmount"]),
        format_inr(profit_data["totalOtherExpenses"]),
        format_inr(profit_data["netProfit"]),
        f"{profit_data['profitMarginPercent']}%"
    ]
    
    ws.append([])
    
    header_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    bold_font = Font(name="Arial", bold=True)
    
    ws.append(summary_headers[:8])
    for col_num in range(1, 9):
        cell = ws.cell(row=3, column=col_num)
        cell.font = bold_font
        cell.fill = header_fill
        
    ws.append(summary_values)
    ws.append([])
    
    trips_headers = [
        "Date", "Trip ID", "Driver ID", "Vehicle ID", "Weight (MT)", 
        "Rate (INR)", "Rate Type", "Revenue (INR)", "Fuel Cost (INR)", 
        "Toll (INR)", "Site Charges (INR)", "Other Expenses (INR)", 
        "Profit (INR)", "Start Location", "End Location", "Notes"
    ]
    ws.append(trips_headers)
    
    header_row = 6
    trips_header_fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
    white_bold = Font(name="Arial", bold=True, color="FFFFFF")
    
    for col_num in range(1, len(trips_headers) + 1):
        cell = ws.cell(row=header_row, column=col_num)
        cell.font = white_bold
        cell.fill = trips_header_fill
        cell.alignment = Alignment(horizontal="center")
        
    for t in sorted(trips, key=lambda x: x.get("date", ""), reverse=True):
        weight_val = float(t.get("materialWeight")) if t.get("materialWeight") is not None else ""
        rate_val = float(t.get("rate")) if t.get("rate") is not None else ""
        rev_val = format_inr(float(t.get("earnings"))) if t.get("earnings") is not None else "Pending"
        fuel_val = format_inr(float(t.get("fuelCost"))) if t.get("fuelCost") is not None else "None"
        profit_val = format_inr(float(t.get("profit"))) if t.get("profit") is not None else "Pending"
        
        row = [
            t.get("date", ""),
            t.get("id", ""),
            t.get("driverId", ""),
            t.get("vehicleId", ""),
            weight_val,
            rate_val,
            t.get("rateType", ""),
            rev_val,
            fuel_val,
            float(t.get("tollAmount", 0.0) or 0.0),
            float(t.get("loadingUnloadingAmount", 0.0) or 0.0),
            float(t.get("otherExpenses", 0.0) or 0.0),
            profit_val,
            t.get("startLocation", ""),
            t.get("endLocation", ""),
            t.get("notes", "")
        ]
        ws.append(row)
        
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
        
    buffer = io.BytesIO()
    wb.save(buffer)
    excel_bytes = buffer.getvalue()
    buffer.close()
    return excel_bytes
